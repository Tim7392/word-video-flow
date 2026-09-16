"""Canonical GUI-independent parsing and five-track export logic.

Moved verbatim from the repaired GUI; frozen AST and byte-output tests protect
these functions. Importing this module creates no window or persistent log.
"""
import sys, os, re, math, logging, tempfile, shutil
from pathlib import Path
from datetime import datetime
try:
    import docx
except ImportError:
    docx = None
_repair_logger = logging.getLogger('subtitle_factory')
_repair_logger.addHandler(logging.NullHandler())

class LogicCore:
    def __init__(self):
        self.re_pos = re.compile(r"\b(n\.|v\.|vt\.|vi\.|adj\.|adv\.|prep\.|conj\.|pron\.|num\.|art\.|a\.)",
                                 re.IGNORECASE)

    def parse(self, path=None, txt=None):
        raw = []
        if txt:
            raw = txt.split('\n')
        elif path:
            try:
                if path.lower().endswith('.docx'):
                    if docx:
                        d = docx.Document(path)
                        raw = [x.text for x in d.paragraphs if x.text.strip()]
                    else:
                        return [], ["请安装 python-docx 以读取Word"]
                else:
                    try:
                        with open(path, 'r', encoding='utf-8') as f:
                            raw = f.readlines()
                    except:
                        with open(path, 'r', encoding='gbk') as f:
                            raw = f.readlines()
            except Exception as e:
                return [], [str(e)]

        data, errs = [], []
        for i, l in enumerate(raw):
            l = l.strip()
            if not l or l.isdigit(): continue

            l = re.sub(r'^\s*\(?\d+(\.\s*|\s+|\、|\))\s*', '', l)
            l = re.sub(r'^[^a-zA-Z]+', '', l)

            m = re.match(r"^([a-zA-Z\-\'\s]+)(.*)", l)
            if not m: continue
            w, rem = m.groups()
            w, rem = w.strip(), rem.strip()

            phone, defin = "", rem
            pm = re.search(r"(\[.*?\]|/.*?/)", rem)
            if pm:
                phone = pm.group(1)
                defin = rem.replace(phone, "").strip()

            clean = self.re_pos.sub('', defin).strip()
            clean = re.sub(r"\[.*?\]|/.*?/", "", clean).strip()  # 二次清理
            data.append({"w": w, "p": phone, "d_f": defin, "d_c": clean})

        return data, errs

    def gen(self, data, out, t1, t2, prefix=""):
        if not data: return
        fs = [[] for _ in range(5)]
        names = ["_01_英文重复.srt", "_02_英文单次.srt", "_03_音标.srt", "_04_中文带词性.srt", "_05_中文无词性.srt"]

        cur = 0.0
        for i, d in enumerate(data):
            seq = i + 1
            n = len(re.findall(r'[\u4e00-\u9fa5]', d['d_c']))
            if n == 0: n = len(d['w']) // 4 if d['w'] else 1
            if n == 0: n = 1

            dt = n * t1 if n <= 6 else (6 * t1 + (n - 6) * t2)
            dt = max(dt, 0.5)

            t0, t1_tm, t2_tm, te = cur, cur + 1, cur + 2, cur + 2 + dt

            def fmt(s):
                h, r = divmod(int(s), 3600);
                m, sc = divmod(r, 60);
                ms = int((s % 1) * 1000)
                return f"{h:02}:{m:02}:{sc:02},{ms:03}"

            tm0, tm1, tm2, tme = fmt(t0), fmt(t1_tm), fmt(t2_tm), fmt(te)

            fs[0].append(f"{seq}\n{tm0} --> {tm1}\n{d['w']}\n\n{seq}\n{tm1} --> {tm2}\n{d['w']}\n\n")
            fs[1].append(f"{seq}\n{tm0} --> {tme}\n{d['w']}\n\n")
            fs[2].append(f"{seq}\n{tm1} --> {tme}\n{d['p']}\n\n")
            fs[3].append(f"{seq}\n{tm2} --> {tme}\n{d['d_f']}\n\n")
            fs[4].append(f"{seq}\n{tm2} --> {tme}\n{d['d_c']}\n\n")
            cur = te

        bn = re.sub(r'[^\w]', '', data[0]['w']) or "Output"
        import codecs
        if prefix: bn = bn + prefix
        for idx, content in enumerate(fs):
            with codecs.open(os.path.join(out, bn + names[idx]), 'w', 'utf-8-sig') as f:
                f.writelines(content)

_pos_pattern = re.compile(r'(?<![a-zA-Z])(?:vt|vi|adj|adv|prep|conj|pron|num|art|n|v|a)\.', re.I)

_phone_pattern = re.compile(r'\[[^\]]*\]|/[^/\n]+/')

_time_pattern = re.compile(r'^\s*\d{1,3}:\d{2}:\d{2}[,.]\d{3}\s*-->\s*\d{1,3}:\d{2}:\d{2}[,.]\d{3}.*$')

def _read_source(path):
    source = Path(path)
    if not source.is_file():
        raise ValueError('输入文件不存在，请重新选择文件。')
    if source.suffix.lower() == '.docx':
        if docx is None:
            raise ValueError('无法读取 Word 文档。请另存为 TXT 后导入。')
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        document = docx.Document(str(source))
        lines = []
        for item in document.iter_inner_content():
            if isinstance(item, Paragraph):
                lines.extend(item.text.splitlines())
            elif isinstance(item, Table):
                for row in item.rows:
                    seen = set()
                    cells = []
                    for cell in row.cells:
                        if cell._tc not in seen:
                            seen.add(cell._tc)
                            cells.append(cell.text.strip())
                    lines.append(' '.join(cells))
        return '\n'.join(lines)
    if source.suffix.lower() not in ('.txt', '.srt'):
        raise ValueError('支持 TXT、DOCX 和单词词表 SRT；其他文件请先转换格式。')
    raw = source.read_bytes()
    encodings = ('utf-16',) if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else ('utf-8-sig', 'gb18030')
    for encoding in encodings:
        try:
            text = raw.decode(encoding)
            if '\x00' in text:
                raise ValueError('文件包含无法识别的内容，请另存为 UTF-8 文本。')
            return text
        except UnicodeDecodeError:
            pass
    raise ValueError('无法识别文件编码，请另存为 UTF-8 文本。')

def _source_lines(text):
    lines = text.lstrip('\ufeff').splitlines()
    if not any(_time_pattern.match(line) for line in lines):
        return list(enumerate(lines, 1))
    # One vocabulary entry may span word / phonetic / meaning lines in a cue.
    result, pending = [], []
    for number, line in enumerate(lines, 1):
        if _time_pattern.match(line) or not line.strip() or line.strip().isdigit():
            if pending:
                result.append((pending[0][0], ' '.join(value.strip() for _, value in pending)))
                pending = []
            continue
        pending.append((number, re.sub(r'<[^>]+>', '', line)))
    if pending:
        result.append((pending[0][0], ' '.join(value.strip() for _, value in pending)))
    return result

def _parse(self, path=None, txt=None):
    try:
        text = txt if txt is not None else _read_source(path) if path else ''
        data, errors = [], []
        for number, original in _source_lines(text):
            line = original.strip()
            if not line or line.isdigit():
                continue
            line = re.sub(r'^\s*[（(]?\d+(?:[.、)）]\s*|\s+)', '', line)
            line = re.sub(r'^[\s•*·\-]+', '', line)
            phone_match = _phone_pattern.search(line)
            pos_match = _pos_pattern.search(line)
            boundaries = [match.start() for match in (phone_match, pos_match) if match]
            non_word = re.search(r"[^a-zA-Z\-'’\s]", line)
            if non_word:
                boundaries.append(non_word.start())
            split_at = min(boundaries) if boundaries else len(line)
            word, remainder = line[:split_at].strip(), line[split_at:].strip()
            if not re.fullmatch(r"[a-zA-Z][a-zA-Z\-'’]*(?:\s+[a-zA-Z][a-zA-Z\-'’]*)*", word):
                errors.append(f'第 {number} 行：{original[:100]}')
                continue
            phone_match = _phone_pattern.search(remainder)
            phone = phone_match.group(0) if phone_match else ''
            full = (remainder[:phone_match.start()] + remainder[phone_match.end():]).strip() if phone_match else remainder
            clean = _pos_pattern.sub('', full).strip()
            data.append({'w': word, 'p': phone, 'd_f': full, 'd_c': clean})
        return data, errors
    except Exception as exc:
        _repair_logger.exception('Input could not be read')
        return [], ['读取失败：' + str(exc)]

_suffixes = ('_01_英文重复.srt', '_02_英文单次.srt', '_03_音标.srt', '_04_中文带词性.srt', '_05_中文无词性.srt')

def _format_time(milliseconds):
    seconds, ms = divmod(milliseconds, 1000)
    minutes, seconds = divmod(seconds, 60)
    hours, minutes = divmod(minutes, 60)
    return f'{hours:02}:{minutes:02}:{seconds:02},{ms:03}'

def _gen(self, data, out, t1, t2, prefix=''):
    if not data:
        raise ValueError('没有有效词条，无法生成。')
    tracks = [[] for _ in _suffixes]
    current = 0
    for index, entry in enumerate(data, 1):
        count = len(re.findall(r'[\u4e00-\u9fa5]', entry['d_c']))
        if count == 0:
            count = max(len(entry['w']) // 4, 1)
        duration = count * t1 if count <= 6 else 6 * t1 + (count - 6) * t2
        # Store time as integer milliseconds so 2.8 seconds cannot become 2.799.
        end = current + 2000 + round(max(duration, .5) * 1000)
        starts = (current, current + 1000, current + 2000)
        def cue(sequence, start, stop, content):
            return f'{sequence}\n{_format_time(start)} --> {_format_time(stop)}\n{content}\n\n'
        tracks[0].append(cue(index * 2 - 1, starts[0], starts[1], entry['w']))
        tracks[0].append(cue(index * 2, starts[1], starts[2], entry['w']))
        tracks[1].append(cue(index, starts[0], end, entry['w']))
        tracks[2].append(cue(index, starts[1], end, entry['p']))
        tracks[3].append(cue(index, starts[2], end, entry['d_f']))
        tracks[4].append(cue(index, starts[2], end, entry['d_c']))
        current = end
    base = (re.sub(r'[^\w]', '', data[0]['w']) or 'Output') + prefix
    for suffix, track in zip(_suffixes, tracks):
        with open(Path(out) / (base + suffix), 'w', encoding='utf-8-sig', newline='\n') as stream:
            stream.writelines(track)

def _generate_packages(core, selected, output, t1, t2, batch_size):
    if not selected:
        raise ValueError('没有有效词条，无法生成。')
    if not all(math.isfinite(value) and value > 0 for value in (t1, t2)):
        raise ValueError('字幕停留时间必须大于 0。')
    output = output.strip().strip('"').strip()
    if not output:
        raise ValueError('请选择输出文件夹。')
    destination = Path(output).expanduser().resolve()
    destination.mkdir(parents=True, exist_ok=True)
    size = batch_size or len(selected)
    packages = [selected[i:i+size] for i in range(0, len(selected), size)]
    expected = []
    for index, package in enumerate(packages, 1):
        prefix = f'_Part{index}' if batch_size else ''
        base = (re.sub(r'[^\w]', '', package[0]['w']) or 'Output') + prefix
        expected.extend(base + suffix for suffix in _suffixes)
    # Preserve earlier exports: use a new run folder when names already exist.
    if any((destination / name).exists() for name in expected):
        destination = Path(tempfile.mkdtemp(prefix=datetime.now().strftime('生成_%Y%m%d_%H%M%S_'), dir=destination))
    published = []
    try:
        with tempfile.TemporaryDirectory(prefix='.subtitle_factory_', dir=destination) as staging:
            for index, package in enumerate(packages, 1):
                core.gen(package, staging, t1, t2, f'_Part{index}' if batch_size else '')
            actual = sorted(path.name for path in Path(staging).iterdir())
            if actual != sorted(expected):
                raise RuntimeError('生成文件数量与预期不一致，请查看错误日志。')
            for name in expected:
                target = destination / name
                # Exclusive create also prevents overwrites from concurrent runs.
                with target.open('xb') as writer:
                    published.append(target)
                    writer.write((Path(staging) / name).read_bytes())
    except Exception:
        for path in published:
            try:
                path.unlink()
            except OSError:
                _repair_logger.exception('Could not remove incomplete export')
        raise
    return destination, len(packages), len(expected)

LogicCore.parse = _parse
LogicCore.gen = _gen
LogicCore.gen_srt = _gen
